"""Tests for `app.ingestion.office_extraction.extract_text` -- the shared
Office/PDF/Excel text extraction module originally written inside
`connectors/sharepoint.py`, pulled out once `connectors/confluence.py`
needed the exact same capability for attachment content. These tests build
real, minimal files in-memory with the same library used to parse them
(`pypdf`/`python-docx`/`openpyxl`), so they exercise the actual parsing
libraries, not a mock -- connector-level tests (`test_sharepoint.py`,
`test_confluence.py`) monkeypatch this module instead, since they only need
to confirm "the connector calls this with the right filename/bytes and uses
the result," not re-verify the parsing itself.
"""

from __future__ import annotations

from io import BytesIO

import docx
from openpyxl import Workbook
from pypdf import PdfWriter

from app.ingestion.office_extraction import extract_text


def _pdf_bytes(text: str) -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    # `pypdf`'s writer has no simple "add a text run" helper -- a blank
    # page's `extract_text()` legitimately returns "", so the PDF test
    # below asserts the extraction *pipeline* runs and returns a string
    # without raising, not specific extracted text.
    del text
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_extract_text_docx_joins_paragraphs() -> None:
    raw_bytes = _docx_bytes(["Restart the checkout service.", "Then clear the queue."])

    assert extract_text("runbook.docx", raw_bytes) == (
        "Restart the checkout service.\n\nThen clear the queue."
    )


def test_extract_text_xlsx_joins_nonempty_cells() -> None:
    raw_bytes = _xlsx_bytes([["Step", "Action"], ["1", "Restart the checkout service."]])

    assert extract_text("runbook.xlsx", raw_bytes) == "Step\tAction\n1\tRestart the checkout service."


def test_extract_text_pdf_runs_without_raising() -> None:
    raw_bytes = _pdf_bytes("Restart the checkout service.")

    # See `_pdf_bytes`'s own comment -- a blank test page extracts to "",
    # not the literal text; this asserts the PDF extraction path runs
    # end-to-end without raising, contrasted by the corrupt-file test below.
    result = extract_text("runbook.pdf", raw_bytes)
    assert isinstance(result, str)


def test_extract_text_unsupported_extension_returns_none() -> None:
    assert extract_text("photo.jpg", b"some bytes") is None


def test_extract_text_corrupt_docx_returns_none_not_raise() -> None:
    # Not a real .docx (zip/XML) file -- python-docx must raise internally,
    # and `extract_text` must swallow that, not propagate it.
    assert extract_text("runbook.docx", b"not a real docx file") is None


def test_extract_text_corrupt_pdf_returns_none_not_raise() -> None:
    assert extract_text("runbook.pdf", b"not a real pdf file") is None


def test_extract_text_corrupt_xlsx_returns_none_not_raise() -> None:
    assert extract_text("runbook.xlsx", b"not a real xlsx file") is None
